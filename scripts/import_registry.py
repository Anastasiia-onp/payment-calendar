# -*- coding: utf-8 -*-
"""
Импорт обязательств из произвольной Excel-выгрузки в формат реестра.

Что делает:
  1. Находит строку заголовков, даже если выше идёт шапка отчёта.
  2. Сопоставляет колонки выгрузки с полями реестра по словарю синонимов.
  3. Разбирает суммы («1 120 000,00 руб.») и даты (текстом или датой).
  4. Отбраковывает строки без суммы или без даты — ничего не додумывает.
  5. Пишет результат в новый файл, исходную выгрузку не меняет.

Выученные соответствия колонок сохраняются в scripts/import_mapping.json
и applyются при следующем запуске — сопоставлять заново не нужно.

Запуск:
    python scripts/import_registry.py 01_input/Выгрузка_Счета_к_оплате_ДЕМО.xlsx
    python scripts/import_registry.py <файл> --type Поступление
    python scripts/import_registry.py <файл> --map "Поставщик=контрагент"
"""

import json
import os
import re
import sys
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

MAPPING_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "import_mapping.json")

# поле реестра -> варианты названий колонок в выгрузках
SYNONYMS = {
    "контрагент": [
        "контрагент", "поставщик", "получатель", "плательщик", "покупатель",
        "наименование контрагента", "кому", "от кого", "организация",
    ],
    "основание": [
        "основание", "документ", "документ-основание", "счет", "счёт",
        "номер документа", "назначение платежа", "назначение", "договор",
    ],
    "проект": [
        "проект", "объект", "объект строительства", "заказ", "подразделение",
        "направление", "цфо",
    ],
    "срок": [
        "срок оплаты", "дата платежа", "дата оплаты", "оплатить до", "срок",
        "плановая дата", "дата", "дата документа",
    ],
    "сумма": [
        "сумма", "сумма к оплате", "к оплате", "сумма документа", "сумма платежа",
        "всего", "итого", "сумма, руб", "сумма руб",
    ],
    "статья": [
        "статья", "статья ддс", "статья расходов", "вид расхода", "вид операции",
        "статья затрат", "категория",
    ],
    "комментарий": [
        "комментарий", "примечание", "описание", "содержание",
    ],
}

REQUIRED = ("контрагент", "срок", "сумма")
MAX_HEADER_SCAN = 20     # в скольких первых строках искать заголовки


class Problem(object):
    def __init__(self, row, text, amount=None):
        self.row = row
        self.text = text
        self.amount = amount or 0.0

    def __str__(self):
        return "строка {r}: {t}".format(r=self.row, t=self.text)


# ---------------------------------------------------------------- разбор значений

def norm(text):
    """Нормализация заголовка: нижний регистр, без лишних знаков."""
    if text is None:
        return ""
    s = str(text).strip().lower().replace("ё", "е")
    s = re.sub(r"[\n\r]+", " ", s)
    s = re.sub(r"[.,:;()]+", " ", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def parse_amount(value):
    """«1 120 000,00 руб.» -> 1120000.0. Возвращает None, если это не сумма."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value)
    s = re.sub(r"[^\d,.\-]", "", s.replace("\xa0", " "))
    if not s:
        return None
    # запятая как десятичный разделитель, точка как разделитель групп
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


DATE_PATTERNS = ["%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"]


def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for pattern in DATE_PATTERNS:
        try:
            return datetime.strptime(s, pattern).date()
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------- сопоставление колонок

def load_learned():
    if not os.path.exists(MAPPING_FILE):
        return {}
    try:
        with open(MAPPING_FILE, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except (ValueError, IOError):
        return {}


def save_learned(learned):
    with open(MAPPING_FILE, "w", encoding="utf-8") as fh:
        json.dump(learned, fh, ensure_ascii=False, indent=2, sort_keys=True)


def guess_field(header, learned):
    """Поле реестра для заголовка колонки, иначе None."""
    key = norm(header)
    if not key:
        return None
    if key in learned:
        return learned[key]
    for field, variants in SYNONYMS.items():
        if key in [norm(v) for v in variants]:
            return field
    # частичное совпадение: «сумма к оплате, руб» -> «сумма к оплате»
    for field, variants in SYNONYMS.items():
        for v in variants:
            nv = norm(v)
            if len(nv) >= 5 and (key.startswith(nv) or nv in key):
                return field
    return None


def find_header_row(ws, learned):
    """Строка заголовков — та, где распознано больше всего полей."""
    best_row, best_score, best_map = None, 0, {}
    limit = min(ws.max_row, MAX_HEADER_SCAN)
    for r in range(1, limit + 1):
        mapping, score = {}, 0
        for c in range(1, ws.max_column + 1):
            field = guess_field(ws.cell(row=r, column=c).value, learned)
            if field and field not in mapping:
                mapping[field] = c
                score += 1
        if score > best_score:
            best_row, best_score, best_map = r, score, mapping
    return best_row, best_map


# ---------------------------------------------------------------- импорт

def import_file(path, default_type="Платёж", extra_map=None, learn=True):
    """
    Возвращает (rows, problems, info).

    rows — строки в формате листа «Реестр».
    """
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    learned = load_learned()
    header_row, mapping = find_header_row(ws, learned)

    if extra_map:
        for header, field in extra_map.items():
            for c in range(1, ws.max_column + 1):
                if norm(ws.cell(row=header_row or 1, column=c).value) == norm(header):
                    mapping[field] = c
                    learned[norm(header)] = field

    info = {
        "sheet": ws.title,
        "header_row": header_row,
        "mapping": {},
        "unmapped": [],
        "control_total": None,   # итог, заявленный в самой выгрузке
    }

    if header_row is None:
        return [], [Problem(0, "не удалось найти строку заголовков")], info

    for field, col in mapping.items():
        info["mapping"][field] = ws.cell(row=header_row, column=col).value
    for c in range(1, ws.max_column + 1):
        header = ws.cell(row=header_row, column=c).value
        if header and c not in mapping.values():
            info["unmapped"].append(str(header))

    missing = [f for f in REQUIRED if f not in mapping]
    if missing:
        return [], [Problem(header_row,
                            "не найдены обязательные колонки: " + ", ".join(missing))], info

    if learn:
        for field, col in mapping.items():
            header = norm(ws.cell(row=header_row, column=col).value)
            if header:
                learned[header] = field
        save_learned(learned)

    rows, problems = [], []

    def cell(r, field):
        col = mapping.get(field)
        return ws.cell(row=r, column=col).value if col else None

    for r in range(header_row + 1, ws.max_row + 1):
        contractor = cell(r, "контрагент")
        amount_raw = cell(r, "сумма")
        due_raw = cell(r, "срок")

        if not contractor and amount_raw is None and due_raw is None:
            continue

        # строка итога выгрузки — не обязательство, а контрольная сумма
        text = str(contractor or "").strip().lower()
        if text.startswith(("итого", "всего", "total")):
            control = parse_amount(amount_raw)
            if control is None:
                for c in range(1, ws.max_column + 1):
                    control = parse_amount(ws.cell(row=r, column=c).value)
                    if control:
                        break
            if control:
                info["control_total"] = control
            continue

        amount = parse_amount(amount_raw)
        due = parse_date(due_raw)

        if not contractor:
            problems.append(Problem(r, "нет контрагента — строка пропущена"))
            continue
        if amount is None:
            problems.append(Problem(
                r, "нет суммы у «{c}» — строка пропущена, данных нет".format(c=contractor)))
            continue
        if amount <= 0:
            problems.append(Problem(
                r, "сумма не положительная у «{c}»".format(c=contractor)))
            continue
        if due is None:
            problems.append(Problem(
                r, "нет срока оплаты у «{c}» — строка пропущена, данных нет".format(
                    c=contractor), amount))
            continue

        rows.append((
            default_type,
            str(cell(r, "проект") or "").strip() or None,
            str(contractor).strip(),
            str(cell(r, "основание") or "").strip() or None,
            str(cell(r, "статья") or "").strip() or None,
            amount,
            due,
            "План",
            str(cell(r, "комментарий") or "").strip() or None,
        ))

    rows.sort(key=lambda x: (x[6], x[2]))
    return rows, problems, info


# ---------------------------------------------------------------- запись результата

REG_HEADERS = ["Тип", "Проект", "Контрагент", "Основание", "Статья ДДС",
               "Сумма, ₽", "Срок платежа", "Статус", "Комментарий"]


def write_rows(rows, out_path, source_name):
    wb = Workbook()
    ws = wb.active
    ws.title = "Импорт"

    ws.cell(row=1, column=1, value="Импортировано из: " + source_name).font = \
        __import__("openpyxl").styles.Font(bold=True)
    ws.cell(row=2, column=1,
            value="Дата импорта: " + date.today().strftime("%d.%m.%Y"))

    for i, name in enumerate(REG_HEADERS, start=1):
        ws.cell(row=4, column=i, value=name).font = \
            __import__("openpyxl").styles.Font(bold=True)

    for r, item in enumerate(rows, start=5):
        for c, value in enumerate(item, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c == 6:
                cell.number_format = '#,##0.00\\ "₽"'
            elif c == 7:
                cell.number_format = "DD.MM.YYYY"

    for col, width in zip("ABCDEFGHI", [14, 26, 24, 32, 22, 16, 15, 14, 30]):
        ws.column_dimensions[col].width = width

    wb.save(out_path)


# ---------------------------------------------------------------- точка входа

def main():
    args = [a for a in sys.argv[1:]]
    if not args:
        print("Укажите файл выгрузки:")
        print("  python scripts/import_registry.py 01_input/Выгрузка_Счета_к_оплате_ДЕМО.xlsx")
        return 1

    path = args[0]
    default_type = "Платёж"
    extra_map = {}

    i = 1
    while i < len(args):
        if args[i] == "--type" and i + 1 < len(args):
            default_type = args[i + 1]
            i += 2
        elif args[i] == "--map" and i + 1 < len(args):
            pair = args[i + 1]
            if "=" in pair:
                header, field = pair.split("=", 1)
                extra_map[header.strip()] = field.strip()
            i += 2
        else:
            i += 1

    rows, problems, info = import_file(path, default_type, extra_map)

    print("Файл:   " + os.path.basename(path))
    print("Лист:   {s}, заголовки в строке {r}".format(
        s=info["sheet"], r=info["header_row"]))
    print()
    print("Сопоставление колонок:")
    for field in ["контрагент", "основание", "проект", "срок", "сумма",
                  "статья", "комментарий"]:
        source = info["mapping"].get(field)
        mark = "  " if source else "! "
        print("  {m}{f:<14} <- {s}".format(
            m=mark, f=field, s=source if source else "не найдено"))
    if info["unmapped"]:
        print("  не использованы: " + ", ".join(info["unmapped"]))
    print()

    if rows:
        print("Импортировано строк: {n}".format(n=len(rows)))
        print()
        print("{:<12} {:<20} {:<30} {:>16}".format("Дата", "Контрагент", "Основание", "Сумма"))
        for item in rows:
            print("{:<12} {:<20} {:<30} {:>16}".format(
                item[6].strftime("%d.%m.%Y"), str(item[2])[:20], str(item[3] or "")[:30],
                format(item[5], ",.2f").replace(",", " ")))
        print()
        print("Итого: {s} ₽".format(
            s=format(sum(r[5] for r in rows), ",.2f").replace(",", " ")))

        base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        out_dir = os.path.join(base, "02_calc")
        os.makedirs(out_dir, exist_ok=True)
        out_path = os.path.join(
            out_dir, "Импорт_{d}.xlsx".format(d=date.today().strftime("%Y-%m-%d")))
        write_rows(rows, out_path, os.path.basename(path))
        print("Результат: 02_calc/" + os.path.basename(out_path))

    if problems:
        print()
        print("Не импортировано ({n}):".format(n=len(problems)))
        for p in problems:
            print("  " + str(p))

    # контрольная сверка с итогом самой выгрузки
    control = info.get("control_total")
    if control:
        imported = sum(r[5] for r in rows)
        skipped = sum(p.amount for p in problems)
        diff = control - imported - skipped
        money = lambda x: format(x, ",.2f").replace(",", " ")
        print()
        print("Сверка с итогом выгрузки:")
        print("  Итог в выгрузке:      {s} ₽".format(s=money(control)))
        print("  Импортировано:        {s} ₽".format(s=money(imported)))
        print("  Пропущено с суммой:   {s} ₽".format(s=money(skipped)))
        if abs(diff) < 0.01:
            print("  Расхождение:          нет, всё учтено")
        else:
            print("  РАСХОЖДЕНИЕ:          {s} ₽ — проверьте выгрузку".format(
                s=money(diff)))

    return 0


if __name__ == "__main__":
    sys.exit(main())
